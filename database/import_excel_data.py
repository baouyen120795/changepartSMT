import os
import glob
import re
import datetime
import sqlite3
import openpyxl

DB_PATH = os.path.join(os.path.dirname(__file__), 'smt_management.db')

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn

def clean_val(v):
    if v is None:
        return ''
    s = str(v).strip()
    return '' if s in ['-', 'N/A', 'None', '- '] else s

def parse_date(v):
    if not v:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, (int, float)):
        try:
            d = datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))
            return d.strftime('%Y-%m-%d 00:00:00')
        except:
            return None
    s = str(v).strip()
    # Check if string matches date
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%Y.%m.%d'):
        try:
            dt = datetime.datetime.strptime(s[:19], fmt)
            return dt.strftime('%Y-%m-%d %H:%M:%S')
        except:
            pass
    return None

def normalize_line(line_str):
    if not line_str:
        return 'Line 1'
    s = str(line_str).upper()
    if 'PM ROOM' in s:
        return 'PM ROOM'
    match = re.search(r'([1-6])', s)
    if match:
        return f"Line {match.group(1)}"
    return 'Line 1'

def infer_part_category(name, desc):
    text = (str(name) + ' ' + str(desc)).lower()
    if any(k in text for k in ['belt', 'curoa', 'đai']):
        return 'Belt'
    if any(k in text for k in ['nozzle', 'đầu hút']):
        return 'Nozzle'
    if any(k in text for k in ['valve', 'van', 'solenoid']):
        return 'Valve'
    if any(k in text for k in ['sensor', 'cảm biến', 'quang']):
        return 'Sensor'
    if any(k in text for k in ['filter', 'lọc']):
        return 'Filter'
    if any(k in text for k in ['bearing', 'vòng bi', 'bi']):
        return 'Bearing'
    if any(k in text for k in ['pump', 'bơm']):
        return 'Pump'
    if any(k in text for k in ['motor', 'động cơ']):
        return 'Motor'
    if any(k in text for k in ['xích', 'chain']):
        return 'Chain'
    if any(k in text for k in ['shaft', 'trục']):
        return 'Shaft'
    if any(k in text for k in ['cylinder', 'xi lanh', 'pittong', 'spring']):
        return 'Mechanical'
    if any(k in text for k in ['tụ', 'ups', 'nguồn', 'điện', 'dây điện', 'pin']):
        return 'Electrical'
    return 'General'

def find_header_row(rows):
    for r_idx, row in enumerate(rows[:10]):
        row_str = ' '.join(clean_val(c).lower() for c in row)
        keywords = ['line', 'machine', 'part', 'pic', 'no', 'date', 'seri']
        matches = sum(1 for k in keywords if k in row_str)
        if matches >= 2:
            return r_idx, row
    return None, None

def import_all_excel():
    conn = get_db()
    cursor = conn.cursor()

    print("=== BẮT ĐẦU QUÉT VÀ NHẬP DỮ LIỆU TỪ TẤT CẢ FILE EXCEL ===")

    # Cache existing equipments and spare_parts
    equip_cache = {} # serial or code -> id
    for row in cursor.execute("SELECT id, equipment_code, serial_number FROM equipments"):
        if row['equipment_code']:
            equip_cache[row['equipment_code'].upper()] = row['id']
        if row['serial_number']:
            equip_cache[row['serial_number'].upper()] = row['id']

    part_cache = {} # part_number -> id
    for row in cursor.execute("SELECT id, part_number FROM spare_parts"):
        if row['part_number']:
            part_cache[row['part_number'].upper()] = row['id']

    total_logs_imported = 0
    total_equipments_added = 0
    total_parts_added = 0

    category_folders = [
        ('1. PM SMT - MOUNTER', 'MOUNTER', 'MNT'),
        ('2. PM SMT - PRINTER', 'PRINTER', 'PRT'),
        ('3. PM SMT - LASER MARKING', 'LASER_MARKING', 'LAS'),
        ('4. PM SMT - BONDING', 'BONDING', 'BND'),
        ('5. PM SMT - REFLOW', 'REFLOW', 'RFL'),
        ('6. PM SMT - LOADER,BUFFER,SLIDER,CONVEYOR,FLIP', 'BUFFER_CONVEYOR', 'BUF'),
        ('7. PM SMT - AUTO LABELT', 'AUTO_LABEL', 'LBL')
    ]

    for folder_name, cat, code_prefix in category_folders:
        folder_path = os.path.join('.', folder_name)
        if not os.path.exists(folder_path):
            continue

        xlsx_files = glob.glob(os.path.join(folder_path, '*.xlsx'))
        for fpath in xlsx_files:
            fname = os.path.basename(fpath)
            if fname.startswith('~$'):
                continue

            try:
                wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
            except Exception as e:
                print(f"[Bỏ qua] Không đọc được {fname}: {e}")
                continue

            # Identify target sheets
            target_sheets = []
            for sname in wb.sheetnames:
                s_lower = sname.lower()
                if any(w in s_lower for w in ['history', 'change part', 'loader', 'auto label', 'laser', 'bonding', 'printer-2026', 'gas spring', 'vacuum pump', 'change new pin']):
                    target_sheets.append(sname)
                elif 'sheet1' in s_lower and cat in ['PRINTER', 'BONDING']:
                    target_sheets.append(sname)

            if not target_sheets:
                # Default to first sheet if only 1 sheet
                if len(wb.sheetnames) == 1:
                    target_sheets = [wb.sheetnames[0]]

            for sname in target_sheets:
                ws = wb[sname]
                rows = list(ws.iter_rows(values_only=True))
                if len(rows) < 3:
                    continue

                header_idx, header_row = find_header_row(rows)
                if header_idx is None:
                    continue

                # Map column indices
                col_map = {}
                for idx, col in enumerate(header_row):
                    if col is None:
                        continue
                    c_clean = str(col).lower().replace('\n', ' ').strip()
                    if 'line' in c_clean and 'col_line' not in col_map:
                        col_map['col_line'] = idx
                    elif any(k in c_clean for k in ['lane']) and 'col_lane' not in col_map:
                        col_map['col_lane'] = idx
                    elif any(k in c_clean for k in ['side', 'bot/top', 't/b']) and 'col_side' not in col_map:
                        col_map['col_side'] = idx
                    elif any(k in c_clean for k in ['machine  no', 'machine no', 'mach no']) and 'col_mno' not in col_map:
                        col_map['col_mno'] = idx
                    elif 'maker' in c_clean and 'col_maker' not in col_map:
                        col_map['col_maker'] = idx
                    elif any(k in c_clean for k in ['type', 'model machine', 'model', 'machine']) and 'col_model' not in col_map:
                        col_map['col_model'] = idx
                    elif any(k in c_clean for k in ['serial', 'seri']) and 'col_serial' not in col_map:
                        col_map['col_serial'] = idx
                    elif any(k in c_clean for k in ['date change', 'date error', 'date']) and 'col_date' not in col_map:
                        col_map['col_date'] = idx
                    elif any(k in c_clean for k in ['part name', 'part number', 'part']) and 'col_part' not in col_map:
                        col_map['col_part'] = idx
                    elif any(k in c_clean for k in ['ghi chú', 'item', 'desc']) and 'col_desc' not in col_map:
                        col_map['col_desc'] = idx
                    elif any(k in c_clean for k in ['q-ty', 'q.ty', "q'ty", 'qty']) and 'col_qty' not in col_map:
                        col_map['col_qty'] = idx
                    elif any(k in c_clean for k in ['pic', 'technician']) and 'col_pic' not in col_map:
                        col_map['col_pic'] = idx
                    elif 'status' in c_clean and 'col_status' not in col_map:
                        col_map['col_status'] = idx

                # Process data rows
                imported_this_sheet = 0
                for r_idx in range(header_idx + 1, len(rows)):
                    r = rows[r_idx]
                    if not any(r):
                        continue

                    raw_line = clean_val(r[col_map.get('col_line', 1)]) if 'col_line' in col_map else ''
                    line_name = normalize_line(raw_line)

                    maker = clean_val(r[col_map.get('col_maker', -1)]) if 'col_maker' in col_map else ''
                    model = clean_val(r[col_map.get('col_model', -1)]) if 'col_model' in col_map else ''
                    serial = clean_val(r[col_map.get('col_serial', -1)]) if 'col_serial' in col_map else ''
                    mach_no = clean_val(r[col_map.get('col_mno', -1)]) if 'col_mno' in col_map else ''
                    side = clean_val(r[col_map.get('col_side', -1)]) if 'col_side' in col_map else ''

                    part_code = clean_val(r[col_map.get('col_part', -1)]) if 'col_part' in col_map else ''
                    part_desc = clean_val(r[col_map.get('col_desc', -1)]) if 'col_desc' in col_map else ''

                    raw_date = r[col_map.get('col_date', -1)] if 'col_date' in col_map else None
                    rep_date = parse_date(raw_date)

                    raw_qty = clean_val(r[col_map.get('col_qty', -1)]) if 'col_qty' in col_map else '1'
                    try:
                        qty = int(re.search(r'\d+', raw_qty).group(0)) if re.search(r'\d+', raw_qty) else 1
                    except:
                        qty = 1

                    pic = clean_val(r[col_map.get('col_pic', -1)]) if 'col_pic' in col_map else 'KTV SMT'
                    if not pic:
                        pic = 'KTV SMT'

                    # Require at least part_code or part_desc to be a valid maintenance record
                    if not part_code and not part_desc:
                        continue

                    # Determine part number and part name
                    if not part_code:
                        part_num = re.sub(r'[^A-Za-z0-9\-\.]', '', part_desc.split('\n')[0])[:30]
                        if not part_num:
                            part_num = f"PART-{cat[:3]}-{len(part_cache)+1}"
                        part_name = part_desc.replace('\n', ' ')
                    else:
                        part_num = part_code
                        part_name = part_desc.replace('\n', ' ') if part_desc else part_code

                    # Find or create equipment
                    eq_key = (serial if serial else f"{code_prefix}-{line_name.replace(' ', '')}-{mach_no or '01'}").upper()
                    equip_id = equip_cache.get(eq_key)

                    if not equip_id:
                        # Construct a unique equipment code
                        eq_code = f"{code_prefix}-{line_name.replace(' ', 'L')}-{mach_no or len(equip_cache)+1}"
                        if eq_code.upper() in equip_cache:
                            eq_code = f"{eq_code}-{len(equip_cache)+1}"

                        eq_name = f"{maker} {model or cat}".strip()
                        if not eq_name or eq_name == cat:
                            eq_name = f"{cat.replace('_', ' ')} System"

                        cursor.execute("""
                            INSERT INTO equipments (equipment_code, equipment_name, category, line_name, model, serial_number, status)
                            VALUES (?, ?, ?, ?, ?, ?, 'Running')
                        """, (eq_code, eq_name, cat, line_name, model, serial))
                        equip_id = cursor.lastrowid
                        equip_cache[eq_key] = equip_id
                        equip_cache[eq_code.upper()] = equip_id
                        total_equipments_added += 1

                    # Find or create spare part
                    part_id = part_cache.get(part_num.upper())
                    if not part_id:
                        p_cat = infer_part_category(part_num, part_name)
                        cursor.execute("""
                            INSERT INTO spare_parts (part_number, part_name, category, stock_quantity, min_stock_alert)
                            VALUES (?, ?, ?, ?, ?)
                        """, (part_num, part_name[:100], p_cat, 15, 5))
                        part_id = cursor.lastrowid
                        part_cache[part_num.upper()] = part_id
                        total_parts_added += 1

                    # Format note / failure reason
                    notes_list = []
                    if side: notes_list.append(f"Side: {side}")
                    if raw_line: notes_list.append(f"Vị trí: {raw_line}")
                    if mach_no: notes_list.append(f"Máy #{mach_no}")
                    notes_str = ' | '.join(notes_list)

                    reason = part_desc if part_desc and part_desc != part_name else f"Thay thế phụ tùng {fname}"

                    if not rep_date:
                        rep_date = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

                    # Insert maintenance log
                    cursor.execute("""
                        INSERT INTO maintenance_logs (equipment_id, part_id, quantity_used, replaced_date, failure_reason, technician_name, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                    """, (equip_id, part_id, qty, rep_date, reason[:200], pic[:100], notes_str))

                    total_logs_imported += 1
                    imported_this_sheet += 1

                if imported_this_sheet > 0:
                    print(f"  + [{fname} -> {sname}]: Đã nhập {imported_this_sheet} bản ghi bảo dưỡng")

            wb.close()

    conn.commit()
    conn.close()

    print("\n========================================================")
    print(f"THÀNH CÔNG! ĐÃ TẢI DỮ LIỆU TỪ TẤT CẢ FILE EXCEL:")
    print(f"- Số bản ghi lịch sử bảo dưỡng đã thêm: {total_logs_imported}")
    print(f"- Số thiết bị mới tạo: {total_equipments_added}")
    print(f"- Số linh kiện kho mới tạo: {total_parts_added}")
    print("========================================================")

if __name__ == '__main__':
    import_all_excel()
